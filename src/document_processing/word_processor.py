# SPDX-FileCopyrightText: 2026 Malte Dreyer
# SPDX-License-Identifier: MIT
"""
Word document processor using python-docx.
Table extraction with merged cells and header detection; inline formatting,
hyperlinks and list types are preserved as markdown.
"""

import logging
import time
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple, Set
import re
import zipfile

try:
    from docx import Document as DocxDocument
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph
    from docx.opc.exceptions import PackageNotFoundError
    from docx.shared import Length
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False
    logging.warning("python-docx not installed. Word processing will not be available.")

from .base_processor import DocumentProcessor
from ..models.document import (
    MarkdownDocument,
    DocumentElement,
    ElementType,
    ProcessingOptions,
    ProcessingResult,
    DocumentMetadata
)

logger = logging.getLogger(__name__)

_W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
_R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
_W = '{%s}' % _W_NS

# Containers whose runs are not visible text of the paragraph: deleted or
# moved-away text (tracked changes) and nested paragraphs (text boxes), which
# are not part of this paragraph's text.
_SKIP_RUN_ANCESTORS = {_W + 'del', _W + 'moveFrom', _W + 'p', _W + 'txbxContent'}

# Numbering formats that mean "bullet" rather than "ordered".
_BULLET_FORMATS = {'bullet', 'none'}


def _hyperlink_target(paragraph, hyperlink) -> Optional[str]:
    """External URL of a w:hyperlink element (None for internal anchors)."""
    rel_id = hyperlink.get('{%s}id' % _R_NS)
    if not rel_id:
        return None
    try:
        rel = paragraph.part.rels[rel_id]
        return rel.target_ref if rel.is_external else None
    except KeyError:
        return None


def iter_text_runs(paragraph):
    """
    Yield (run, hyperlink_url) for every visible text run of a paragraph in
    document order: direct runs and runs inside hyperlinks, tracked
    insertions, fields and inline content controls. Deleted text and text
    boxes are skipped. hyperlink_url is the external link target or None.
    """
    from docx.text.run import Run

    p = paragraph._p
    for r in p.iter(_W + 'r'):
        url = None
        skip = False
        ancestor = r.getparent()
        while ancestor is not None and ancestor is not p:
            if ancestor.tag in _SKIP_RUN_ANCESTORS:
                skip = True
                break
            if ancestor.tag == _W + 'hyperlink' and url is None:
                url = _hyperlink_target(paragraph, ancestor)
            ancestor = ancestor.getparent()
        if not skip:
            yield Run(r, paragraph), url


def paragraph_visible_text(paragraph) -> str:
    """Visible text of a paragraph (see iter_text_runs)."""
    return ''.join(run.text for run, _ in iter_text_runs(paragraph))


class CellInfo:
    """Information about a table cell including merging and formatting."""

    def __init__(self, cell: '_Cell', row_idx: int, col_idx: int):
        self.cell = cell
        self.row_idx = row_idx
        self.col_idx = col_idx
        self.text = self._extract_text()
        self.is_merged = self._check_if_merged()
        self.is_header = self._detect_header()
        self.rowspan, self.colspan = self._get_span()
        self.formatting = self._extract_formatting()

    def _extract_text(self) -> str:
        """Extract clean text from cell."""
        try:
            text_parts = []
            for paragraph in self.cell.paragraphs:
                para_text = paragraph_visible_text(paragraph).strip()
                if para_text:
                    text_parts.append(para_text)
            return '\n'.join(text_parts)
        except Exception:
            return ""

    def _check_if_merged(self) -> bool:
        """Check if cell is part of a merged cell."""
        try:
            # Check if cell has merge properties
            tc = self.cell._tc
            return tc.get('gridSpan') is not None or tc.get('vMerge') is not None
        except Exception:
            return False

    def _detect_header(self) -> bool:
        """
        Detect if a cell looks like a header.

        Only formatting is used, deliberately NOT text patterns: an earlier
        version treated any single capitalised word as a header, so data cells
        containing names or cities ('Anna', 'Berlin') were classified as
        headers. Together with the multi-header merging in
        get_headers_and_data() that silently collapsed every row of such a
        table into the header line and dropped all data.
        """
        try:
            bold_runs = 0
            total_runs = 0

            for paragraph in self.cell.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    total_runs += 1
                    if run.bold:
                        bold_runs += 1

            return total_runs > 0 and (bold_runs / total_runs) > 0.7

        except Exception:
            return False

    def _get_span(self) -> Tuple[int, int]:
        """Get rowspan and colspan of cell."""
        try:
            tc = self.cell._tc

            # Get colspan (gridSpan)
            colspan = 1
            grid_span = tc.get('gridSpan')
            if grid_span is not None:
                colspan = int(grid_span)

            # Get rowspan (vMerge)
            rowspan = 1
            # Note: Determining rowspan from vMerge is complex and requires
            # analyzing multiple cells; a simple approach is used.

            return rowspan, colspan
        except Exception:
            return 1, 1

    def _extract_formatting(self) -> Dict[str, Any]:
        """Extract formatting information from cell."""
        formatting = {
            'bold': False,
            'italic': False,
            'alignment': 'left',
            'background_color': None
        }

        try:
            # Check text formatting
            bold_count = 0
            italic_count = 0
            total_runs = 0

            for paragraph in self.cell.paragraphs:
                # Check paragraph alignment
                if paragraph.alignment:
                    alignment_map = {
                        WD_PARAGRAPH_ALIGNMENT.CENTER: 'center',
                        WD_PARAGRAPH_ALIGNMENT.RIGHT: 'right',
                        WD_PARAGRAPH_ALIGNMENT.JUSTIFY: 'justify'
                    }
                    formatting['alignment'] = alignment_map.get(paragraph.alignment, 'left')

                for run in paragraph.runs:
                    total_runs += 1
                    if run.bold:
                        bold_count += 1
                    if run.italic:
                        italic_count += 1

            if total_runs > 0:
                formatting['bold'] = (bold_count / total_runs) > 0.5
                formatting['italic'] = (italic_count / total_runs) > 0.5

        except Exception as e:
            logger.debug(f"Error extracting cell formatting: {e}")

        return formatting


class TableStructureAnalyzer:
    """Analyzes Word table structure to detect headers and merged cells."""

    def __init__(self, table: Table):
        self.table = table
        self.cell_map = self._build_cell_map()
        self.header_rows = self._detect_header_rows()
        self.merged_regions = self._identify_merged_regions()

    def _build_cell_map(self) -> Dict[Tuple[int, int], CellInfo]:
        """Build a map of all cells with their information."""
        cell_map = {}

        for row_idx, row in enumerate(self.table.rows):
            for col_idx, cell in enumerate(row.cells):
                cell_info = CellInfo(cell, row_idx, col_idx)
                cell_map[(row_idx, col_idx)] = cell_info

        return cell_map

    def _detect_header_rows(self) -> Set[int]:
        """
        Determine which rows are header rows.

        Row 0 is always the header (a markdown table requires one, and
        get_headers_and_data() falls back to row 0 anyway). Additional header
        rows are only accepted when Word explicitly marks them as repeating
        header rows (w:tblHeader); guessing them from text patterns swallows
        entire tables (every capitalised single word would count as header).
        """
        header_rows = set()

        if not self.table.rows:
            return header_rows

        header_rows.add(0)

        for row_idx in range(1, min(3, len(self.table.rows))):
            if self._is_repeat_header_row(row_idx):
                header_rows.add(row_idx)
                logger.debug(f"Row {row_idx} marked as repeating header row (w:tblHeader)")

        return header_rows

    def _is_repeat_header_row(self, row_idx: int) -> bool:
        """True if Word marks this row as a repeating header row."""
        try:
            tr = self.table.rows[row_idx]._tr
            tr_pr = tr.find(
                '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}trPr'
            )
            if tr_pr is None:
                return False
            marker = tr_pr.find(
                '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader'
            )
            return marker is not None
        except Exception:
            return False

    def _identify_merged_regions(self) -> List[Dict[str, Any]]:
        """Identify merged cell regions."""
        merged_regions = []

        # This is a simplified approach - full implementation would be more complex
        for (row_idx, col_idx), cell_info in self.cell_map.items():
            if cell_info.colspan > 1 or cell_info.rowspan > 1:
                merged_regions.append({
                    'start_row': row_idx,
                    'start_col': col_idx,
                    'end_row': row_idx + cell_info.rowspan - 1,
                    'end_col': col_idx + cell_info.colspan - 1,
                    'text': cell_info.text
                })

        return merged_regions

    def get_effective_grid(self) -> List[List[str]]:
        """Get effective grid handling merged cells."""
        if not self.table.rows:
            return []

        max_cols = max(len(row.cells) for row in self.table.rows)
        grid = []

        for row_idx, row in enumerate(self.table.rows):
            row_data = []
            for col_idx in range(max_cols):
                cell_info = self.cell_map.get((row_idx, col_idx))
                if cell_info:
                    row_data.append(cell_info.text)
                else:
                    row_data.append("")  # Empty cell or merged cell continuation
            grid.append(row_data)

        return grid

    def get_headers_and_data(self) -> Tuple[List[str], List[List[str]]]:
        """Get separated headers and data rows."""
        grid = self.get_effective_grid()

        if not grid:
            return [], []

        headers = []
        data_rows = []

        # Extract headers
        if self.header_rows:
            # Use first header row as primary headers
            primary_header_row = min(self.header_rows)
            headers = grid[primary_header_row]

            # If there are multiple header rows, combine them
            if len(self.header_rows) > 1:
                combined_headers = []
                for col_idx in range(len(headers)):
                    header_parts = []
                    for header_row in sorted(self.header_rows):
                        if col_idx < len(grid[header_row]):
                            part = grid[header_row][col_idx].strip()
                            if part:
                                header_parts.append(part)
                    combined_headers.append(' '.join(header_parts))
                headers = combined_headers
        else:
            # No headers detected, use first row
            headers = grid[0] if grid else []

        # Extract data rows (skip header rows)
        skip_rows = self.header_rows if self.header_rows else {0}
        for row_idx, row in enumerate(grid):
            if row_idx not in skip_rows:
                data_rows.append(row)

        # Clean headers
        headers = [str(h).strip() for h in headers]

        return headers, data_rows


class WordProcessor(DocumentProcessor):
    """
    Processor for Word documents with robust table handling.
    """

    def __init__(self, options: Optional[ProcessingOptions] = None):
        """Initialize Word processor."""
        super().__init__(options)
        if not PYTHON_DOCX_AVAILABLE:
            raise ImportError("python-docx is required for Word processing. Install with: pip install python-docx")

        # KISS: Only map KNOWN styles, everything else defaults to PARAGRAPH
        self.heading_styles = {
            'Title': ElementType.HEADING1,
            'Heading 1': ElementType.HEADING1,
            'Heading 2': ElementType.HEADING2,
            'Heading 3': ElementType.HEADING3,
            'Heading 4': ElementType.HEADING4,
            'Heading 5': ElementType.HEADING5,
            'Heading 6': ElementType.HEADING6,
            'Überschrift 1': ElementType.HEADING1,  # German
            'Überschrift 2': ElementType.HEADING2,
            'Überschrift 3': ElementType.HEADING3,
        }

        self.list_styles = {
            'List Paragraph',
            'List Bullet',
            'List Number',
            'Liste',
            'Aufzählung',
        }

        self.quote_styles = {
            'Quote',
            'Intense Quote',
            'Zitat',
        }

    def can_process(self, file_path: Path) -> bool:
        """Check if file is a processable Word document."""
        if not PYTHON_DOCX_AVAILABLE:
            return False

        suffix = file_path.suffix.lower()

        if suffix == '.docx':
            return self._is_valid_docx(file_path)
        elif suffix == '.doc':
            return True

        return False

    def get_supported_formats(self) -> List:
        """Return supported formats."""
        from . import DocumentFormat  # Local import to avoid circular dependency
        return [DocumentFormat.DOCX, DocumentFormat.DOC]

    def _is_valid_docx(self, file_path: Path) -> bool:
        """Check if file is a valid .docx file."""
        try:
            with zipfile.ZipFile(file_path, 'r') as zip_file:
                required_files = ['[Content_Types].xml', 'word/document.xml']
                zip_contents = set(zip_file.namelist())

                for required_file in required_files:
                    if required_file not in zip_contents:
                        return False

                return True

        except (zipfile.BadZipFile, FileNotFoundError, PermissionError):
            return False

    def _detect_actual_format(self, file_path: Path) -> str:
        """Detect the actual file format regardless of extension."""
        try:
            with open(file_path, 'rb') as f:
                header = f.read(8)

            if header[:4] == b'PK\x03\x04':
                if self._is_valid_docx(file_path):
                    return 'docx'
                else:
                    return 'zip_corrupted'
            elif header[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                return 'doc_legacy'
            elif header[:4] == b'%PDF':
                return 'pdf'
            elif header.startswith(b'\xff\xfe') or header.startswith(b'\xfe\xff'):
                return 'text_unicode'

            return 'unknown'

        except Exception as e:
            logger.warning(f"Could not detect file format: {e}")
            return 'unknown'

    async def validate_file(self, file_path: Path) -> Tuple[bool, Optional[str]]:
        """File validation with format detection."""
        is_valid, error_msg = await super().validate_file(file_path)
        if not is_valid:
            return False, error_msg

        actual_format = self._detect_actual_format(file_path)
        file_extension = file_path.suffix.lower()

        if file_extension == '.docx':
            if actual_format == 'docx':
                return True, None
            elif actual_format == 'doc_legacy':
                return False, (
                    "File has .docx extension but is actually a legacy .doc file. "
                    "Please save as .docx format in Word or use conversion tools."
                )
            elif actual_format in ['zip_corrupted', 'unknown']:
                return False, (
                    "File appears to be corrupted or is not a valid Word document. "
                    "Try re-saving the document in Word."
                )
            else:
                return False, f"File is not a Word document (detected: {actual_format})"

        elif file_extension == '.doc':
            if actual_format == 'doc_legacy':
                return False, (
                    "Legacy .doc format detected. This system only supports .docx files. "
                    "Please convert to .docx format:\n"
                    "1. Open in Microsoft Word\n"
                    "2. Save As → Word Document (.docx)\n" 
                    "3. Upload the .docx version"
                )
            elif actual_format == 'docx':
                return False, (
                    "File is actually in .docx format but has .doc extension. "
                    "Please rename to .docx extension."
                )
            else:
                return False, f"File is not a Word document (detected: {actual_format})"

        return True, None

    async def process(self, file_path: Path) -> ProcessingResult:
        """
        Process a Word file with robust table handling.
        """
        start_time = time.time()

        # Validation
        is_valid, error_msg = await self.validate_file(file_path)
        if not is_valid:
            return ProcessingResult.error_result(error_msg)

        try:
            if not self._is_valid_docx(file_path):
                return ProcessingResult.error_result(
                    "File validation failed. Please ensure this is a valid .docx file."
                )

            # Open Word document
            try:
                doc = DocxDocument(str(file_path))
            except PackageNotFoundError as e:
                return ProcessingResult.error_result(
                    "Unable to open Word document. The file may be corrupted or password-protected."
                )
            except ValueError as e:
                error_str = str(e)
                if "themeManager" in error_str:
                    return ProcessingResult.error_result(
                        "Document appears to be corrupted or contains invalid Office components."
                    )
                else:
                    return ProcessingResult.error_result(f"Word document error: {error_str}")

            # DEBUG: Analyze document structure
            self._debug_document_structure(doc)

            # Create metadata
            metadata = self._extract_metadata(doc, file_path)

            # Extract content with enhanced table processing
            elements = self._extract_content_enhanced(doc)

            # Minimal post-processing
            elements = self._minimal_post_processing(elements)

            if not elements:
                return ProcessingResult.error_result(
                    f"No readable content found in document. "
                    f"Document has {len(doc.paragraphs)} paragraphs detected."
                )

            # Rebuild raw text
            raw_text_parts = [elem.content for elem in elements]
            raw_text = '\n'.join(raw_text_parts)

            # Create markdown content from elements
            markdown_parts = []
            for element in elements:
                if hasattr(element, 'to_markdown'):
                    markdown_parts.append(element.to_markdown())
                else:
                    markdown_parts.append(element.content)

            markdown_content = '\n\n'.join(markdown_parts)

            # Create document with required content parameter
            document = MarkdownDocument(
                content=markdown_content,  # required parameter
                elements=elements,
                metadata=metadata,
                raw_text=raw_text,
                processing_time=time.time() - start_time
            )

            # Calculate statistics
            stats = self.calculate_statistics(document)
            logger.info(f"Word processing completed: {stats}")

            return ProcessingResult.success_result(
                document=document,
                processing_time=time.time() - start_time
            )

        except Exception as e:
            logger.error(f"Word processing error: {str(e)}", exc_info=True)
            return ProcessingResult.error_result(f"Word processing failed: {str(e)}")

    def _extract_content_enhanced(self, doc: DocxDocument) -> List[DocumentElement]:
        """
        Extract content in document order with robust table processing.
        """
        logger.info("=== CONTENT EXTRACTION ===")

        elements = []

        # Track position in document
        position = 0

        # Map body XML elements to their python-docx objects once. Searching
        # the paragraph/table lists for every body element made the import
        # quadratic in the document length.
        paragraphs_by_element = {para._element: para for para in doc.paragraphs}
        tables_by_element = {table._element: table for table in doc.tables}

        # Process paragraphs and tables in document order
        for element in doc.element.body:
            if element.tag == _W + 'p':  # Paragraph
                para = paragraphs_by_element.get(element)
                if para is not None:
                    para_element = self._process_paragraph_kiss(para, position)
                    if para_element:
                        elements.append(para_element)
                        position += 1

            elif element.tag == _W + 'tbl':  # Table
                table = tables_by_element.get(element)
                if table is not None:
                    table_elements = self._process_table_enhanced(table, position)
                    if table_elements:
                        elements.extend(table_elements)
                        position += len(table_elements)

        logger.info(f"Enhanced extraction: {len(elements)} elements extracted")
        return elements

    def _process_table_enhanced(self, table: Table, position: int) -> List[DocumentElement]:
        """
        Process a Word table with structure analysis.
        """
        logger.info(f"=== PROCESSING TABLE AT POSITION {position} ===")

        try:
            if not table.rows:
                logger.warning("Table has no rows")
                return []

            # Analyze table structure
            analyzer = TableStructureAnalyzer(table)

            logger.info(f"Table analysis: {len(table.rows)} rows, "
                       f"header rows: {analyzer.header_rows}, "
                       f"merged regions: {len(analyzer.merged_regions)}")

            # Get structured data
            headers, data_rows = analyzer.get_headers_and_data()

            if not headers:
                logger.warning("No headers found in table")
                return []

            # Create enhanced table data using the local TableData class
            from . import TableData  # Local import to avoid circular dependency
            table_data = TableData(headers=headers, rows=data_rows)

            # Generate markdown representation
            markdown_content = table_data.to_markdown()

            # Validate the generated markdown
            from .markdown_converter import TableValidator
            if not TableValidator.is_valid_markdown_table(markdown_content):
                logger.warning("Generated invalid markdown table, attempting repair")
                from .markdown_converter import TableReconstructor
                repaired_content, was_repaired = TableReconstructor.repair_table(markdown_content)
                if was_repaired and TableValidator.is_valid_markdown_table(repaired_content):
                    markdown_content = repaired_content
                    logger.info("Table markdown repaired successfully")
                else:
                    logger.error("Could not repair table markdown")

            # Create table element with enhanced metadata
            table_element = DocumentElement(
                type=ElementType.TABLE,
                content=markdown_content,
                metadata={
                    'rows': len(data_rows) + len(analyzer.header_rows),
                    'columns': len(headers),
                    'has_header': bool(analyzer.header_rows),
                    'header_rows': list(analyzer.header_rows),
                    'merged_cells': len(analyzer.merged_regions),
                    'cell_count': len(analyzer.cell_map),
                    'table_analysis': {
                        'structure_complexity': len(analyzer.merged_regions),
                        'formatting_preserved': True
                    }
                },
                position=position
            )

            logger.info(f"Table processed successfully: {len(headers)} columns, "
                       f"{len(data_rows)} data rows, {len(analyzer.header_rows)} header rows")

            return [table_element]

        except Exception as e:
            logger.error(f"Enhanced table processing failed: {e}", exc_info=True)

            # Fallback to simple table processing
            logger.info("Falling back to simple table processing")
            return self._process_table_fallback(table, position)

    def _process_table_fallback(self, table: Table, position: int) -> List[DocumentElement]:
        """Fallback table processing for when enhanced processing fails."""
        try:
            if not table.rows:
                return []

            # Simple extraction - treat first row as headers
            first_row = table.rows[0]
            headers = []
            for cell in first_row.cells:
                cell_text = ""
                for paragraph in cell.paragraphs:
                    para_text = paragraph.text.strip()
                    if para_text:
                        cell_text += para_text + " "
                headers.append(self.clean_text(cell_text.strip()))

            # Extract data rows
            data_rows = []
            for row in table.rows[1:]:
                row_data = []
                for cell in row.cells:
                    cell_text = ""
                    for paragraph in cell.paragraphs:
                        para_text = paragraph.text.strip()
                        if para_text:
                            cell_text += para_text + " "
                    row_data.append(self.clean_text(cell_text.strip()))

                # Ensure row has same number of columns as headers
                while len(row_data) < len(headers):
                    row_data.append("")
                if len(row_data) > len(headers):
                    row_data = row_data[:len(headers)]

                data_rows.append(row_data)

            # Create table data using local TableData class
            from . import TableData  # Local import to avoid circular dependency
            table_data = TableData(headers=headers, rows=data_rows)

            # Create table element
            table_element = DocumentElement(
                type=ElementType.TABLE,
                content=table_data.to_markdown(),
                metadata={
                    'rows': len(data_rows) + 1,
                    'columns': len(headers),
                    'has_header': True,
                    'processing_method': 'fallback'
                },
                position=position
            )

            logger.info(f"Fallback table processing successful: {len(headers)} columns, {len(data_rows)} rows")
            return [table_element]

        except Exception as e:
            logger.error(f"Fallback table processing also failed: {e}")

            # Last resort - create text representation
            try:
                table_text = self._table_to_text(table)
                if table_text:
                    return [DocumentElement(
                        type=ElementType.PARAGRAPH,
                        content=f"**Table** (could not preserve structure):\n\n{table_text}",
                        metadata={'failed_table': True},
                        position=position
                    )]
            except Exception as final_e:
                logger.error(f"Even text conversion failed: {final_e}")

            return []

    def _table_to_text(self, table: Table) -> str:
        """Convert table to plain text representation."""
        text_lines = []

        for row_idx, row in enumerate(table.rows):
            row_texts = []
            for cell in row.cells:
                cell_text = ""
                for paragraph in cell.paragraphs:
                    para_text = paragraph.text.strip()
                    if para_text:
                        cell_text += para_text + " "
                row_texts.append(cell_text.strip())

            if any(row_texts):  # Only add non-empty rows
                text_lines.append(" | ".join(row_texts))

        return "\n".join(text_lines)

    def _process_paragraph_kiss(self, paragraph: Paragraph, position: int) -> Optional[DocumentElement]:
        """
        KISS paragraph processing - simple rules, robust fallbacks
        """
        # Text of all visible runs, including runs inside hyperlinks, tracked
        # insertions, fields and inline content controls (paragraph.runs only
        # returns direct children, which silently dropped e.g. link texts).
        text = ''.join(run.text for run, _ in self._iter_text_runs(paragraph)).strip()

        if not text:
            logger.debug(f"Para[{position}] SKIPPED: No text found")
            return None

        # Get style
        style_name = paragraph.style.name if paragraph.style else "Unknown"

        # KISS element type determination
        element_type = self._determine_element_type_kiss(text, style_name, paragraph)

        # Basic metadata
        metadata = {
            'style': style_name,
            'original_position': position
        }

        # Add formatting info
        if paragraph.runs:
            first_run = paragraph.runs[0]
            metadata['bold'] = bool(first_run.bold)
            metadata['italic'] = bool(first_run.italic)

        # For body text, capture inline bold/italic as markdown so that a single
        # bold/italic word inside a sentence survives translation and export.
        # Headings stay plain (they are rendered via '#' and are usually uniform).
        inline_types = (ElementType.PARAGRAPH, ElementType.LIST_ITEM, ElementType.QUOTE)
        if element_type in inline_types:
            inline_md = self._runs_to_inline_markdown(paragraph)
            content_text = self.clean_text(inline_md) if inline_md.strip() else self.clean_text(text)
            # Inline markup is embedded in the content itself, so the
            # converter must not wrap the whole paragraph again.
            metadata['bold'] = False
            metadata['italic'] = False
        else:
            content_text = self.clean_text(text)

        # Create element
        element = DocumentElement(
            type=element_type,
            content=content_text,
            metadata=metadata,
            position=position
        )

        # Set heading level
        if element.type.value.startswith('h'):
            element.level = int(element.type.value[1])

        # Set list metadata
        if element_type == ElementType.LIST_ITEM:
            element.level = self._get_word_list_level(paragraph)
            element.metadata['list_type'] = self._get_word_list_type(paragraph)

        logger.debug(f"Para[{position}] PROCESSED: {element_type.value} (style: {style_name})")
        return element

    def _iter_text_runs(self, paragraph: Paragraph):
        """See iter_text_runs()."""
        return iter_text_runs(paragraph)

    def _runs_to_inline_markdown(self, paragraph: Paragraph) -> str:
        """
        Build inline markdown (**bold**, *italic*, ***both***, [text](url))
        from a paragraph's runs. Adjacent runs with identical formatting and
        link target are merged, and emphasis markers are placed tight against
        the text (markdown does not allow spaces next to the markers).
        """
        merged: List[Tuple[str, bool, bool, Optional[str]]] = []
        for run, url in self._iter_text_runs(paragraph):
            if not run.text:
                continue
            bold = bool(run.bold)
            italic = bool(run.italic)
            if merged and merged[-1][1:] == (bold, italic, url):
                prev_text = merged[-1][0]
                merged[-1] = (prev_text + run.text, bold, italic, url)
            else:
                merged.append((run.text, bold, italic, url))

        parts = []
        for text, bold, italic, url in merged:
            stripped = text.strip()
            if not stripped:
                parts.append(text)
                continue
            lead = text[:len(text) - len(text.lstrip())]
            trail = text[len(text.rstrip()):]
            inner = stripped
            if bold or italic:
                marker = '***' if (bold and italic) else ('**' if bold else '*')
                inner = f"{marker}{inner}{marker}"
            if url:
                inner = f"[{inner}]({url})"
            parts.append(f"{lead}{inner}{trail}")

        return ''.join(parts)

    def _determine_element_type_kiss(self, text: str, style_name: str, paragraph: Paragraph) -> ElementType:
        """
        KISS element type determination: Simple rules, everything else = PARAGRAPH
        """
        # Rule 1: Check known heading styles
        if style_name in self.heading_styles:
            logger.debug(f"Style '{style_name}' → {self.heading_styles[style_name].value}")
            return self.heading_styles[style_name]

        # Rule 2: Check known list styles
        if style_name in self.list_styles:
            logger.debug(f"Style '{style_name}' → LIST_ITEM")
            return ElementType.LIST_ITEM

        # Rule 3: Check known quote styles
        if style_name in self.quote_styles:
            logger.debug(f"Style '{style_name}' → QUOTE")
            return ElementType.QUOTE

        # Rule 4: Check for obvious list patterns (regardless of style)
        if self._has_list_marker(text):
            logger.debug(f"Text pattern → LIST_ITEM")
            return ElementType.LIST_ITEM

        # Rule 5: Check for obvious heading patterns (regardless of style)
        if self._looks_like_heading(text, paragraph):
            logger.debug(f"Text pattern → HEADING4")
            return ElementType.HEADING4

        # Rule 6: DEFAULT = PARAGRAPH
        logger.debug(f"Style '{style_name}' → PARAGRAPH (default)")
        return ElementType.PARAGRAPH

    def _has_list_marker(self, text: str) -> bool:
        """Check if text has obvious list markers"""
        list_patterns = [
            r'^\s*[-*•·▪▫◦‣]\s+',    # Bullets
            r'^\s*\d+[\.\)]\s+',       # Numbers
            r'^\s*[a-zA-Z][\.\)]\s+',  # Letters
            r'^\s*[IVX]+[\.\)]\s+',    # Roman numerals
        ]

        for pattern in list_patterns:
            if re.match(pattern, text):
                return True
        return False

    def _looks_like_heading(self, text: str, paragraph: Paragraph) -> bool:
        """Check if text looks like a heading"""
        # Short, bold text
        if len(text) < 100 and paragraph.runs:
            bold_runs = sum(1 for run in paragraph.runs if run.bold)
            if bold_runs == len(paragraph.runs) and bold_runs > 0:
                return True

        # All caps, short text
        if len(text) < 80 and text.isupper() and len(text.split()) <= 10:
            return True

        # Text ending with colon
        if len(text) < 120 and text.endswith(':') and len(text.split()) <= 15:
            return True

        return False

    def _get_word_list_level(self, paragraph: Paragraph) -> int:
        """
        Get the actual list level from Word structure
        """
        try:
            # Check numbering level (ilvl)
            if hasattr(paragraph, '_element') and paragraph._element is not None:
                p_element = paragraph._element
                if hasattr(p_element, 'pPr') and p_element.pPr is not None:
                    num_pr = p_element.pPr.numPr
                    if num_pr is not None and hasattr(num_pr, 'ilvl') and num_pr.ilvl is not None:
                        level = int(num_pr.ilvl.val) + 1  # Word uses 0-based, we use 1-based
                        logger.debug(f"Word list level: {level}")
                        return level

            # Fallback: estimate from indentation
            if hasattr(paragraph, 'paragraph_format') and paragraph.paragraph_format.left_indent:
                indent_pt = paragraph.paragraph_format.left_indent.pt if hasattr(paragraph.paragraph_format.left_indent, 'pt') else 0
                level = max(1, int(indent_pt / 36) + 1)  # 36pt = 0.5 inch per level
                return level

        except Exception as e:
            logger.debug(f"Error getting Word list level: {e}")

        return 1

    def _get_word_list_type(self, paragraph: Paragraph) -> str:
        """
        'ordered' or 'unordered', taken from the Word numbering definition.

        The numbering can be set on the paragraph or inherited from its style
        (e.g. "List Number"); its format (decimal, lowerLetter, bullet, ...)
        decides. Falls back to the text pattern if no definition is found.
        """
        try:
            num_fmt = self._resolve_numbering_format(paragraph)
            if num_fmt is not None:
                return 'unordered' if num_fmt in _BULLET_FORMATS else 'ordered'
        except Exception as e:
            logger.debug(f"Error getting Word list type: {e}")

        # Fallback to simple text pattern
        return self._detect_list_type_simple(paragraph.text)

    def _resolve_numbering_format(self, paragraph: Paragraph) -> Optional[str]:
        """Return w:numFmt of the paragraph's list level, or None."""
        num_id, ilvl = None, None

        def read_num_pr(ppr):
            if ppr is None:
                return None, None
            num_pr = ppr.find(_W + 'numPr')
            if num_pr is None:
                return None, None
            nid = num_pr.find(_W + 'numId')
            lvl = num_pr.find(_W + 'ilvl')
            return (nid.get(_W + 'val') if nid is not None else None,
                    lvl.get(_W + 'val') if lvl is not None else None)

        num_id, ilvl = read_num_pr(paragraph._p.pPr)
        style = paragraph.style
        while num_id is None and style is not None:
            s_num, s_lvl = read_num_pr(style.element.pPr)
            num_id = s_num
            ilvl = ilvl if ilvl is not None else s_lvl
            style = style.base_style
        if num_id is None or num_id == '0':
            return None

        numbering_part = paragraph.part.numbering_part
        numbering = numbering_part.element
        num = next((n for n in numbering.findall(_W + 'num')
                    if n.get(_W + 'numId') == num_id), None)
        if num is None:
            return None
        abstract_ref = num.find(_W + 'abstractNumId')
        if abstract_ref is None:
            return None
        abstract_id = abstract_ref.get(_W + 'val')
        abstract = next((a for a in numbering.findall(_W + 'abstractNum')
                         if a.get(_W + 'abstractNumId') == abstract_id), None)
        if abstract is None:
            return None
        level = ilvl or '0'
        for lvl in abstract.findall(_W + 'lvl'):
            if lvl.get(_W + 'ilvl') == level:
                fmt = lvl.find(_W + 'numFmt')
                return fmt.get(_W + 'val') if fmt is not None else None
        return None

    def _detect_list_type_simple(self, text: str) -> str:
        """Simple list type detection"""
        if re.match(r'^\s*\d+[\.\)]\s+', text):
            return 'ordered'
        elif re.match(r'^\s*[a-zA-Z][\.\)]\s+', text):
            return 'ordered'
        elif re.match(r'^\s*[IVX]+[\.\)]\s+', text):
            return 'ordered'
        else:
            return 'unordered'

    def _minimal_post_processing(self, elements: List[DocumentElement]) -> List[DocumentElement]:
        """
        Minimal post-processing - only remove truly empty elements
        """
        logger.info("=== MINIMAL POST-PROCESSING ===")

        if not elements:
            return elements

        # Only filter truly empty content
        filtered_elements = []
        for element in elements:
            if element and element.content and element.content.strip():
                filtered_elements.append(element)
            else:
                logger.debug(f"Removing empty element at position {element.position if element else 'unknown'}")

        logger.info(f"Minimal post-processing: {len(elements)} → {len(filtered_elements)} elements")

        # Update positions
        for idx, element in enumerate(filtered_elements):
            element.position = idx

        return filtered_elements

    def _debug_document_structure(self, doc: DocxDocument):
        """
        DEBUG: Document structure analysis with enhanced table info
        """
        logger.info("=== DOCUMENT STRUCTURE ===")
        logger.info(f"Total paragraphs: {len(doc.paragraphs)}")
        logger.info(f"Total tables: {len(doc.tables)}")

        # Analyze tables
        for i, table in enumerate(doc.tables):
            try:
                analyzer = TableStructureAnalyzer(table)
                logger.info(f"Table {i}: {len(table.rows)} rows, "
                           f"header rows: {analyzer.header_rows}, "
                           f"merged regions: {len(analyzer.merged_regions)}")
            except Exception as e:
                logger.warning(f"Could not analyze table {i}: {e}")

        # Count by style
        styles = {}
        unknown_styles = set()

        for p in doc.paragraphs:
            style = p.style.name if p.style else "None"
            styles[style] = styles.get(style, 0) + 1

            # Track unknown styles
            if (style not in self.heading_styles and
                style not in self.list_styles and
                style not in self.quote_styles and
                style not in ['Normal', 'None']):
                unknown_styles.add(style)

        logger.info(f"Style distribution: {styles}")

        if unknown_styles:
            logger.info(f"Unknown styles (will default to PARAGRAPH): {unknown_styles}")

    def _extract_metadata(self, doc: DocxDocument, file_path: Path) -> DocumentMetadata:
        """Extract metadata from Word document."""
        try:
            core_properties = doc.core_properties
            metadata_kwargs = {}

            try:
                if core_properties.title:
                    metadata_kwargs['title'] = core_properties.title
            except:
                pass

            try:
                if core_properties.author:
                    metadata_kwargs['author'] = core_properties.author
            except:
                pass

            metadata = self.create_metadata(file_path, **metadata_kwargs)

            try:
                word_count = sum(len(p.text.split()) for p in doc.paragraphs)
                metadata.word_count = word_count
                metadata.page_count = max(1, word_count // 300)
            except:
                metadata.word_count = 0
                metadata.page_count = 1

            try:
                char_count = sum(len(p.text) for p in doc.paragraphs)
                metadata.char_count = char_count
            except:
                metadata.char_count = 0

            return metadata

        except Exception as e:
            logger.warning(f"Error extracting metadata: {e}")
            return self.create_metadata(file_path)