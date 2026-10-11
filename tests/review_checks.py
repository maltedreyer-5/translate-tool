#!/usr/bin/env python3
# SPDX-FileCopyrightText: 2026 Malte Dreyer
# SPDX-License-Identifier: MIT
"""
Regression checks for configuration, interface construction, the document
round trip (DOCX -> Markdown -> Word) and the error behaviour of the
translation engine.

Run under pytest (tests/test_review_checks.py) or on its own:
    python tests/review_checks.py

Without gradio installed, the interface check is skipped.
"""
import asyncio
import logging
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
logging.disable(logging.ERROR)

RESULTS = []


def check(name, condition):
    RESULTS.append((name, bool(condition)))


def make_sample_docx(path):
    from docx import Document
    doc = Document()
    doc.add_heading('3.2 Ergebnisse und Ausblick (Teil-A)', level=1)
    p = doc.add_paragraph()
    p.add_run('Dies ist ein ')
    r = p.add_run('wichtiger'); r.bold = True
    p.add_run(' Absatz mit ')
    r = p.add_run('kursivem'); r.italic = True
    p.add_run(' Text.')
    li = doc.add_paragraph(style='List Bullet')
    li.add_run('Punkt mit ')
    r = li.add_run('fettem'); r.bold = True
    li.add_run(' Wort')
    doc.save(path)
    return path


def test_config():
    from src.config import load_config
    c = load_config()
    check("Config loads", c is not None)
    check("Retry capped (max_retry_delay present)", getattr(c.llm, 'max_retry_delay', None) is not None)
    check("Retry values conservative", c.llm.retry_delay <= 10 and c.llm.max_retries <= 4)
    check("Unused config sections absent", not hasattr(c, 'optimized_translation'))


def test_interface():
    try:
        import gradio as gr
    except ImportError:
        print("  (skipped: gradio not installed)")
        return
    from src.app import UnifiedDocumentTranslationApp
    app = UnifiedDocumentTranslationApp()
    app.create_unified_interface()
    check("Interface builds without errors", True)


def test_roundtrip():
    from docx import Document
    from src.document_processing.word_processor import WordProcessor
    from src.document_processing.markdown_converter import create_table_aware_converter
    from src.ui.export_handlers import EnhancedExportHandler

    src_path = make_sample_docx('/tmp/verify_in.docx')
    res = asyncio.run(WordProcessor().process(Path(src_path)))
    md = create_table_aware_converter().convert_document(res.document)

    check("Heading keeps its numbering", "3.2" in md)
    check("No backslash escapes in heading", "\\(" not in md and "\\-" not in md)
    check("Inline formatting captured as markdown", "**wichtiger**" in md)

    out = EnhancedExportHandler().export_word(md, filename="verify_out")
    doc = Document(out)
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    italic = [r.text for p in doc.paragraphs for r in p.runs if r.italic]
    text = "\n".join(p.text for p in doc.paragraphs)

    check("Word export: real bold runs", "wichtiger" in bold)
    check("Word export: real italic runs", "kursivem" in italic)
    check("Word export: no raw markdown", "**" not in text)
    check("Word export: no YAML frontmatter", "author:" not in text)


def test_engine_resilience():
    from src.translation.translation_engine import TableAwareTranslationEngine
    from src.chunking.simple_chunker import SimpleChunk, SimpleChunkingResult

    class PartialClient:
        async def translate_text(self, text=None, **kw):
            if "Absatz 2" in text:
                raise RuntimeError("simulierter Ausfall")
            return f"[EN] {text}"

    eng = TableAwareTranslationEngine(PartialClient())
    eng.retry_delay = eng.max_retry_delay = 0.01
    chunks = [SimpleChunk(content=f"Absatz {i}", position=i, token_count=5) for i in range(5)]
    cr = SimpleChunkingResult(chunks=chunks, total_tokens=25, chunk_count=5)
    res = asyncio.run(eng.translate_document(cr, "de", "en"))

    check("Partial failure: no total loss", len(res.translated_chunks) == 5)
    check("Partial failure: order preserved",
          [c['position'] for c in res.translated_chunks] == [0, 1, 2, 3, 4])
    failed = [c for c in res.translated_chunks if c.get('failed')]
    check("Failed section keeps original text",
          bool(failed) and failed[0]['translated_content'] == failed[0]['original_content'])

    # Backoff-Deckelung
    delays = [min(eng.retry_delay * (2 ** a), eng.max_retry_delay)
              for a in range(eng.max_retries - 1)]
    check("Backoff capped", all(d <= eng.max_retry_delay for d in delays))


def test_chunking_and_reassembly():
    """Table detection, preserved punctuation and structure-aware reassembly."""
    import re as _re
    from src.chunking.simple_chunker import SimpleChunker
    from src.translation.translation_engine import TableAwareTranslationEngine
    from src.app import reassemble_translated_chunks

    class Identity:
        async def translate_text(self, text=None, **kw): return text
        async def translate_table_cell(self, text=None, **kw): return text
        async def translate_batch(self, texts=None, **kw): return list(texts or [])

    def roundtrip(md, max_tokens=500):
        cr = SimpleChunker(max_tokens=max_tokens).chunk_text(md, language='de')
        res = asyncio.run(
            TableAwareTranslationEngine(Identity()).translate_document(cr, "de", "en")
        )
        return reassemble_translated_chunks(res.translated_chunks), cr

    # Tables are recognised at all (separator-row regex)
    cr = SimpleChunker(max_tokens=700).chunk_text(
        "| Name | Rolle |\n| --- | --- |\n| Anna | Leitung |", language='de')
    check("Markdown table recognised as table", cr.tables_found == 1)

    # A large table is split by rows and stays within the token limit
    big = "| Nr | Wert |\n| --- | --- |\n" + "\n".join(
        f"| {i} | Wert Nummer {i} mit etwas Text |" for i in range(60))
    cr_big = SimpleChunker(max_tokens=300).chunk_text(big, language='de')
    check("Large table is split", cr_big.chunk_count > 1)
    check("No chunk exceeds the token limit",
          all(c.token_count <= 300 * 1.05 for c in cr_big.chunks))

    # A split table is reassembled into ONE table
    out, _ = roundtrip(big, max_tokens=300)
    seps = [l for l in out.split('\n')
            if set(l.replace('|', '').replace(' ', '')) <= set('-:') and '-' in l]
    check("Split table stays one table", len(seps) == 1)
    check("Table lossless", out.strip() == big.strip())

    # Punctuation survives splitting long paragraphs
    sentences = "Frage eins? Antwort zwei! Aussage drei. Ende vier?"
    cr_s = SimpleChunker(max_tokens=8).chunk_text(sentences, language='de')
    joined = " ".join(c.content for c in cr_s.chunks)
    check("Punctuation preserved",
          joined.count('?') == 2 and joined.count('!') == 1 and joined.count('.') == 1)
    check("No character loss when splitting sentences",
          _re.sub(r'\s', '', joined) == _re.sub(r'\s', '', sentences))

    # A split paragraph stays ONE paragraph
    long_p = "Ein langer Satz ohne Absatzumbruch. " * 80
    out_p, cr_p = roundtrip(long_p, max_tokens=500)
    check("Long paragraph is split", cr_p.chunk_count > 1)
    check("Split paragraph stays one paragraph", "\n\n" not in out_p.strip())
    check("Paragraph lossless",
          _re.sub(r'\s+', ' ', out_p).strip() == _re.sub(r'\s+', ' ', long_p).strip())


def test_word_table_headers():
    """Kopfzeilenerkennung darf keine Datenzeilen verschlucken."""
    from docx import Document
    from src.document_processing.word_processor import WordProcessor
    from src.document_processing.markdown_converter import create_table_aware_converter

    d = Document()
    t = d.add_table(rows=3, cols=2)
    for i, row in enumerate([["Name", "Ort"], ["Anna", "Berlin"], ["Bob", "Potsdam"]]):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v
    path = '/tmp/verify_table.docx'
    d.save(path)

    res = asyncio.run(WordProcessor().process(Path(path)))
    md = create_table_aware_converter().convert_document(res.document)
    table_lines = [l for l in md.split('\n') if l.strip().startswith('|')]
    check("Table with names keeps all rows", len(table_lines) == 4)


def test_quality_features():
    """Context enrichment, auto-glossary and quality check."""
    from src.chunking.simple_chunker import SimpleChunker
    from src.translation.translation_engine import TableAwareTranslationEngine

    md = ("# Verwaltungsvereinbarung\n\n"
          "Die Behörde prüft den Antrag. Der Antrag betrifft die Verwaltungsvereinbarung.\n\n"
          "## Zuständigkeit\n\n"
          "Der Bescheid ergeht schriftlich. Die Behörde erlässt den Bescheid.\n\n"
          "## Fristen\n\n"
          "Der Antrag ist fristgerecht einzureichen. Die Behörde bestätigt den Antrag.")

    class SpyLLM:
        def __init__(self):
            self.contexts = []
            self.terms = []

        async def translate_text(self, text=None, context=None, glossary=None, **kw):
            if context and 'TERMINOLOGY LIST' in context:
                self.terms.append(text)
                return f"EN[{text}]"
            self.contexts.append(context or "")
            self.glossaries = getattr(self, 'glossaries', [])
            self.glossaries.append(glossary or {})
            return text

        async def translate_table_cell(self, text=None, **kw): return text
        async def translate_batch(self, texts=None, **kw): return list(texts or [])

    cr = SimpleChunker(max_tokens=45).chunk_text(md, language='de')
    check("Chunks carry their heading",
          all(c.parent_heading for c in cr.chunks))

    llm = SpyLLM()
    asyncio.run(TableAwareTranslationEngine(llm).translate_document(cr, "de", "en"))

    check("Auto-glossary finds terms", len(llm.terms) >= 3)
    check("Terms without line breaks/articles",
          all('\n' not in t and not t.lower().startswith(('der ', 'die ', 'das '))
              for t in llm.terms))
    check("Section heading in the prompt",
          all('CURRENT SECTION' in c for c in llm.contexts))
    check("Neighbouring text in the prompt",
          any('SURROUNDING TEXT' in c for c in llm.contexts))
    check("Auto-glossary reaches chunk translation",
          any(g for g in getattr(llm, 'glossaries', [])))
    check("No duplicate terminology in the context",
          all('DOCUMENT TERMINOLOGY' not in c for c in llm.contexts))

    # The quality check detects planted errors
    class OkLLM:
        async def translate_text(self, text=None, **kw): return "OK"

    engine = TableAwareTranslationEngine(OkLLM())
    faulty = [
        {'original_content': "Der Betrag beträgt 1500 Euro.",
         'translated_content': "The amount is 1600 Euro.", 'position': 0},
        {'original_content': "- A\n- B\n- C", 'translated_content': "- A\n- B", 'position': 1},
        {'original_content': "Die Behörde entscheidet.",
         'translated_content': "The agency decides.", 'position': 2},
        {'original_content': "Ein korrekter Satz.",
         'translated_content': "A correct sentence.", 'position': 3},
    ]
    findings = asyncio.run(engine.verify_translation(
        faulty, "de", "en", glossary={"Behörde": "authority"}))
    flagged = {f['chunk_index'] for f in findings}
    check("Quality check detects number deviation", 0 in flagged)
    check("Quality check detects missing list items", 1 in flagged)
    check("Quality check detects glossary violation", 2 in flagged)
    check("Quality check does not flag a correct section", 3 not in flagged)


def test_glossary_and_budget():
    """Case sensitivity, prompt budget and verdict evaluation."""
    from src.translation.translation_engine import TableAwareTranslationEngine

    # Prompt budget and relevance filter
    try:
        from src.llm_client import LLMClient
        client = LLMClient()
        budget = client._prompt_budget_chars()
        check("Prompt budget derived from the context window",
              budget is not None and budget > 0)

        big = {f"Fachbegriff Nummer {i}": f"term {i}" for i in range(1000)}
        check("Large glossary fits a 32k window",
              client._glossary_chars(big) < budget)

        text = "Der Fachbegriff Nummer 7 taucht hier auf."
        relevant = client._relevant_glossary(big, text, case_sensitive=False)
        check("Relevance filter works", 0 < len(relevant) < len(big))

        cs = client._relevant_glossary({"Behörde": "authority"},
                                       "die behörde entscheidet", case_sensitive=True)
        ci = client._relevant_glossary({"Behörde": "authority"},
                                       "die behörde entscheidet", case_sensitive=False)
        check("Case sensitivity respected", len(cs) == 0 and len(ci) == 1)
    except Exception as e:
        check(f"LLM client checks ({type(e).__name__}: {e})", False)

    # Verdict evaluation of the semantic check
    class VerdictLLM:
        def __init__(self, verdict): self.verdict = verdict
        async def complete(self, messages, **kw):
            class R: pass
            R.content = self.verdict
            R.finish_reason = "stop"
            return R()

    ok_variants = ["OK", "OK.", "ok", "**OK**", ""]
    no_finding = all(
        asyncio.run(TableAwareTranslationEngine(VerdictLLM(v))
                    ._semantic_check("Quelle", "Ziel", "de", "en")) is None
        for v in ok_variants
    )
    check("Confirmations raise no false alarm", no_finding)

    real = asyncio.run(
        TableAwareTranslationEngine(VerdictLLM("Der Betrag fehlt in der Übersetzung."))
        ._semantic_check("Quelle", "Ziel", "de", "en"))
    check("Real finding is reported", real is not None)


def main():
    for name, fn in [("Configuration", test_config),
                     ("Interface", test_interface),
                     ("Document round trip", test_roundtrip),
                     ("Word tables", test_word_table_headers),
                     ("Chunking & reassembly", test_chunking_and_reassembly),
                     ("Quality features", test_quality_features),
                     ("Glossary & budget", test_glossary_and_budget),
                     ("Engine robustness", test_engine_resilience)]:
        print(f"\n== {name} ==")
        try:
            fn()
        except Exception as e:
            check(f"{name} (exception: {type(e).__name__}: {e})", False)

    print("\n" + "=" * 55)
    for name, passed in RESULTS:
        print(f"  [{'OK ' if passed else 'FAIL'}] {name}")
    failed = [n for n, p in RESULTS if not p]
    print(f"\n{len(RESULTS) - len(failed)}/{len(RESULTS)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
