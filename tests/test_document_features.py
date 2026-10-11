# SPDX-FileCopyrightText: 2026 Malte Dreyer
# SPDX-License-Identifier: MIT
"""
In-place Word translation, model-built glossary, row-wise tables, partial
results and the prompt language. Model answers come from stubs.
"""
import asyncio
import json
import logging
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import RGBColor

from src.config import load_config
from src.llm_client import LLMClient, CompletionResponse
from src.chunking.simple_chunker import SimpleChunker
from src.translation.translation_engine import TableAwareTranslationEngine
from src.translation.docx_inplace import (DocxInPlaceTranslator, encode_paragraph,
                                          apply_translation)

logging.disable(logging.CRITICAL)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def upper_keep_markers(text):
    return re.sub(r'(<[^>]+>)|([^<]+)', lambda m: m.group(1) or m.group(2).upper(), text)


class StubClient(LLMClient):
    """Uppercases the text, keeps markers; optional answer transformation."""

    def __init__(self, config, transform=None, glossary_answer=None):
        super().__init__(config)
        self.transform = transform
        self.glossary_answer = glossary_answer
        self.requests = []

    async def complete(self, messages, **kw):
        system, user = messages[0].content, messages[-1].content
        self.requests.append((system, user))
        if "Answer with JSON only" in system:
            content = self.glossary_answer if self.glossary_answer is not None else "no json"
            return CompletionResponse(content=content, model="stub", finish_reason="stop")
        text = user.split("\n\n", 1)[1] if "\n\n" in user else user
        answer = upper_keep_markers(text)
        if self.transform:
            answer = self.transform(answer)
        return CompletionResponse(content=answer, model="stub", finish_reason="stop")


def config(**translation):
    cfg = load_config()
    cfg.translation.auto_glossary = False
    for key, value in translation.items():
        setattr(cfg.translation, key, value)
    return cfg


def make_docx(path: Path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Kopfzeile"
    doc.sections[0].footer.paragraphs[0].text = "Fußzeile"
    doc.add_heading("Bericht", 1)
    p = doc.add_paragraph("Das ist ")
    bold = p.add_run("wichtig")
    bold.bold = True
    p.add_run(" für alle.")
    q = doc.add_paragraph()
    red = q.add_run("Rot")
    red.font.color.rgb = RGBColor(255, 0, 0)
    q.add_run(" und Zeile")
    q.runs[-1].add_break()
    q.add_run("zwei")
    q.alignment = 1
    table = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Aufgabe", "Status"), ("Schulung", "offen"), ("Bericht", "offen")]):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
    doc.add_paragraph("1234")  # nothing to translate
    doc.save(path)


def translate_docx(tmp_path, cfg, client=None):
    src = tmp_path / "in.docx"
    make_docx(src)
    out = tmp_path / "out.docx"
    client = client or StubClient(cfg)
    engine = TableAwareTranslationEngine(client, cfg)
    result = asyncio.run(DocxInPlaceTranslator(engine, cfg).translate_file(
        src, out, "de", "en", "preserve_original", "formal", {}))
    return Document(out), result, client


# --------------------------------------------------------------------------
# in-place Word translation
# --------------------------------------------------------------------------
def test_inplace_keeps_layout_and_formatting(tmp_path):
    doc, result, _ = translate_docx(tmp_path, config())
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "BERICHT" and doc.paragraphs[0].style.name == "Heading 1"
    assert texts[1] == "DAS IST WICHTIG FÜR ALLE."
    assert [r.text for r in doc.paragraphs[1].runs if r.bold] == ["WICHTIG"]
    red = [r.text for r in doc.paragraphs[2].runs
           if r.font.color and r.font.color.rgb == RGBColor(255, 0, 0)]
    assert red == ["ROT"]
    assert doc.paragraphs[2].alignment == 1
    assert doc.paragraphs[2]._p.findall(f'.//{W}br'), "soft line break lost"
    assert texts[3] == "1234"
    assert doc.sections[0].header.paragraphs[0].text == "KOPFZEILE"
    assert doc.sections[0].footer.paragraphs[0].text == "FUSSZEILE"
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows] == [
        ["AUFGABE", "STATUS"], ["SCHULUNG", "OFFEN"], ["BERICHT", "OFFEN"]]
    assert result.failed_paragraphs == 0 and result.formatting_fallbacks == 0


def test_inplace_identical_cells_answered_from_cache(tmp_path):
    _, result, client = translate_docx(tmp_path, config(table_mode="cell"))
    cell_requests = [u for _, u in client.requests if "offen" in u]
    assert len(cell_requests) == 1


def test_inplace_row_mode_groups_cells(tmp_path):
    doc, _, client = translate_docx(tmp_path, config(table_mode="row"))
    assert any("<p1>Schulung</p1>\n<p2>offen</p2>" in u for _, u in client.requests)
    assert doc.tables[0].rows[1].cells[0].text == "SCHULUNG"


def test_inplace_moved_markers_move_the_runs(tmp_path):
    def swap(answer):
        # "<s1>DAS IST </s1><s2>WICHTIG</s2><s3> FÜR ALLE.</s3>" -> s2 first
        m = re.search(r'(<s1>.*?</s1>)(<s2>.*?</s2>)', answer)
        return answer.replace(m.group(0), m.group(2) + m.group(1)) if m else answer
    cfg = config()
    doc, result, _ = translate_docx(tmp_path, cfg, StubClient(cfg, transform=swap))
    runs = [r for r in doc.paragraphs[1].runs if r.text]
    assert runs[0].text == "WICHTIG" and runs[0].bold
    assert result.formatting_fallbacks == 0


def test_inplace_damaged_markers_fall_back_to_plain_text(tmp_path):
    cfg = config()
    damage = lambda a: a.replace("<s2>", "").replace("</s2>", "")
    doc, result, _ = translate_docx(tmp_path, cfg, StubClient(cfg, transform=damage))
    assert doc.paragraphs[1].text == "DAS IST WICHTIG FÜR ALLE."
    assert result.formatting_fallbacks >= 1


def test_inplace_missing_paragraph_is_retried_alone(tmp_path):
    cfg = config()
    calls = {"n": 0}

    def drop_second(answer):
        calls["n"] += 1
        return re.sub(r'<p2>.*?</p2>', '', answer, flags=re.DOTALL) if "<p2>" in answer and calls["n"] == 1 else answer
    doc, result, _ = translate_docx(tmp_path, cfg, StubClient(cfg, transform=drop_second))
    assert doc.paragraphs[1].text == "DAS IST WICHTIG FÜR ALLE."
    assert result.failed_paragraphs == 0


def test_inplace_footnotes_are_translated(tmp_path):
    import subprocess, shutil
    if not shutil.which("pandoc"):
        pytest.skip("pandoc not available to create a document with footnotes")
    md = tmp_path / "f.md"
    md.write_text("Text mit Fußnote[^1].\n\n[^1]: Die Erklärung steht hier.\n")
    src = tmp_path / "f.docx"
    subprocess.run(["pandoc", str(md), "-o", str(src)], check=True)
    cfg = config()
    engine = TableAwareTranslationEngine(StubClient(cfg), cfg)
    out = tmp_path / "f_out.docx"
    asyncio.run(DocxInPlaceTranslator(engine, cfg).translate_file(
        src, out, "de", "en", "preserve_original", "formal", {}))
    footnotes = zipfile.ZipFile(out).read("word/footnotes.xml").decode()
    assert "DIE ERKLÄRUNG STEHT HIER." in footnotes
    assert 'w:footnoteReference' in zipfile.ZipFile(out).read("word/document.xml").decode()


def test_encode_paragraph_marks_objects():
    doc = Document()
    p = doc.add_paragraph("Bild ")
    run = p.add_run()
    run._r.append(OxmlElement('w:drawing'))
    p.add_run(" danach")
    encoded, _ = encode_paragraph(p._p)
    assert encoded.encoded == "<s1>Bild </s1><x2/><s3> danach</s3>"
    assert apply_translation(encoded, "<s1>Image </s1><x2/><s3> after</s3>")
    assert p._p.find(f'.//{W}drawing') is not None


# --------------------------------------------------------------------------
# glossary built by the model in one request
# --------------------------------------------------------------------------
def _chunks(cfg, text):
    return SimpleChunker(max_tokens=2000, config=cfg).chunk_text(text, "de").chunks


def test_model_glossary_one_request_and_filters():
    cfg = config(auto_glossary_method="model")
    answer = json.dumps({"terms": [
        {"source": "Projektgruppe", "target": "project group"},
        {"source": "Erfunden", "target": "invented"},          # not in the document
        {"source": "Leitung", "target": "a whole sentence " * 10},  # not a term
    ]})
    client = StubClient(cfg, glossary_answer=f"<think>x</think>```json\n{answer}\n```")
    engine = TableAwareTranslationEngine(client, cfg)
    text = "Die Projektgruppe tagt. Die Projektgruppe berichtet der Leitung."
    glossary = asyncio.run(engine.build_document_glossary(_chunks(cfg, text), "de", "en"))
    assert glossary == {"Projektgruppe": "project group"}
    assert len(client.requests) == 1


def test_model_glossary_falls_back_to_heuristic():
    cfg = config(auto_glossary_method="model")
    client = StubClient(cfg, glossary_answer="sorry, no JSON")
    engine = TableAwareTranslationEngine(client, cfg)
    text = "Die Projektgruppe tagt. Die Projektgruppe berichtet."
    glossary = asyncio.run(engine.build_document_glossary(_chunks(cfg, text), "de", "en"))
    assert glossary.get("Projektgruppe") == "PROJEKTGRUPPE"   # translated per term by the stub
    assert len(client.requests) == 2


# --------------------------------------------------------------------------
# row-wise tables in the text pipeline
# --------------------------------------------------------------------------
TABLE = "| Aufgabe | Status |\n|---|---|\n| Schulung | offen |\n| Bericht | erledigt |"


def test_row_mode_translates_rows_in_one_request():
    cfg = config(table_mode="row")
    client = StubClient(cfg)
    engine = TableAwareTranslationEngine(client, cfg)
    chunks = SimpleChunker(max_tokens=2000, config=cfg).chunk_text(TABLE, "de")
    result = asyncio.run(engine.translate_document(chunks, "de", "en", "preserve_original"))
    out = result.translated_chunks[0]['translated_content']
    assert "| SCHULUNG | OFFEN |" in out and "| BERICHT | ERLEDIGT |" in out
    row_requests = [u for _, u in client.requests if "<c1>" in u]
    assert len(row_requests) == 2


def test_row_mode_falls_back_to_cells():
    cfg = config(table_mode="row")
    client = StubClient(cfg, transform=lambda a: re.sub(r'</?c2>', '', a))
    engine = TableAwareTranslationEngine(client, cfg)
    chunks = SimpleChunker(max_tokens=2000, config=cfg).chunk_text(TABLE, "de")
    result = asyncio.run(engine.translate_document(chunks, "de", "en", "preserve_original"))
    assert "| SCHULUNG | OFFEN |" in result.translated_chunks[0]['translated_content']


# --------------------------------------------------------------------------
# partial results and prompt language
# --------------------------------------------------------------------------
def test_finished_sections_are_reported():
    cfg = config()
    engine = TableAwareTranslationEngine(StubClient(cfg), cfg)
    seen = []
    engine.on_chunk_done = lambda index, text: seen.append(index)
    text = "\n\n".join(f"Absatz{i} " + "Wort " * 80 for i in range(3))
    chunks = SimpleChunker(max_tokens=100, config=cfg).chunk_text(text, "de")
    asyncio.run(engine.translate_document(chunks, "de", "en", "preserve_original"))
    assert sorted(seen) == list(range(len(chunks.chunks)))


@pytest.mark.parametrize("language,expected,target_name", [
    ("de", "Du bist ein hochqualifizierter", "Englisch"),
    ("en", "You are a highly qualified", "English"),
])
def test_prompt_language(language, expected, target_name):
    cfg = config(prompt_language=language)
    client = StubClient(cfg)
    asyncio.run(client.translate_text("Hallo", "de", "en", style="akademisch", glossary={"A": "B"}))
    system, user = client.requests[0]
    assert system.startswith(expected) and target_name in system
    assert ("GLOSSARY:" in system) == (language == "en")


def test_prompt_eval_chrf():
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
    from prompt_eval import chrf
    assert chrf("Same text.", "Same text.") == pytest.approx(100.0)
    assert chrf("Other.", "Same text.") < 30
