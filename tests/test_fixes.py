# SPDX-FileCopyrightText: 2026 Malte Dreyer
# SPDX-License-Identifier: MIT
"""
Regression tests for behaviour that is easy to break silently.

No model server is needed: model calls are answered by stubs or by a local
aiohttp test server.
"""
import asyncio
import atexit
import logging
import os
import time
from pathlib import Path

import pytest
from aiohttp import web
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.config import load_config, deep_merge, resolve_path, PROJECT_ROOT
from src.llm_client import (LLMClient, CompletionResponse, Message, MessageRole,
                            LLMTimeoutError)
from src.markdown_utils import (is_table_separator, split_table_row,
                                count_table_columns, is_valid_table_lines)
from src.chunking.simple_chunker import SimpleChunker
from src.translation.translation_engine import TableAwareTranslationEngine
from src.document_processing.word_processor import WordProcessor
from src.document_processing.markdown_converter import create_table_aware_converter
from src.ui.export_handlers import EnhancedExportHandler

logging.disable(logging.CRITICAL)

HYPERLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def add_hyperlink(paragraph, text, url):
    rid = paragraph.part.relate_to(url, HYPERLINK_REL, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def docx_to_markdown(path):
    result = asyncio.run(WordProcessor().process(Path(path)))
    assert result.success, result.error_message
    return create_table_aware_converter(include_frontmatter=False).convert_document(result.document)


class EchoClient(LLMClient):
    """Answers every request with the text it was asked to translate."""

    def __init__(self, config=None):
        super().__init__(config or load_config())
        self.calls = []

    async def complete(self, messages, temperature=None, max_tokens=None, stream=False, **kw):
        user = messages[-1].content
        self.calls.append((messages[0].content, user))
        text = user.split("\n\n", 1)[1] if "\n\n" in user else user
        return CompletionResponse(content=text, model="stub", finish_reason="stop")


# --------------------------------------------------------------------------
# Word import / export
# --------------------------------------------------------------------------
def test_hyperlink_text_survives_import_and_export(tmp_path):
    doc = Document()
    p = doc.add_paragraph("See ")
    add_hyperlink(p, "LINKTEXT", "https://example.org")
    p.add_run(" for details.")
    src = tmp_path / "in.docx"
    doc.save(src)

    md = docx_to_markdown(src)
    assert "[LINKTEXT](https://example.org)" in md

    out = Document(EnhancedExportHandler(temp_dir=tmp_path).export_word(md))
    first = out.paragraphs[0]
    assert "LINKTEXT" in first.text
    assert first._p.findall(qn('w:hyperlink')), "export lost the hyperlink"


def test_tracked_insertion_in_table_cell_is_imported(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Header"
    para = table.cell(1, 0).paragraphs[0]
    para.add_run("Visible")
    ins = OxmlElement('w:ins')
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = " INSERTED"
    run.append(t)
    ins.append(run)
    para._p.append(ins)
    src = tmp_path / "ins.docx"
    doc.save(src)
    assert "INSERTED" in docx_to_markdown(src)


def test_numbered_list_stays_numbered(tmp_path):
    doc = Document()
    doc.add_paragraph("First", style="List Number")
    doc.add_paragraph("Second", style="List Number")
    doc.add_paragraph("Bullet", style="List Bullet")
    src = tmp_path / "list.docx"
    doc.save(src)
    md = docx_to_markdown(src)
    assert "1. First" in md and "1. Second" in md
    assert "- Bullet" in md

    out = Document(EnhancedExportHandler(temp_dir=tmp_path).export_word(md))
    styles = [p.style.name for p in out.paragraphs if p.text]
    assert styles == ["List Number", "List Number", "List Bullet"]


def test_document_metadata_is_not_content(tmp_path):
    doc = Document()
    doc.core_properties.author = "Erika Mustermann"
    doc.core_properties.title = "Internal memo"
    doc.add_paragraph("Hello world.")
    src = tmp_path / "meta.docx"
    doc.save(src)
    md = docx_to_markdown(src)
    assert "Mustermann" not in md and "Internal memo" not in md

    handler = EnhancedExportHandler(temp_dir=tmp_path)
    for export in (handler.export_markdown, handler.export_html, handler.export_json):
        content = Path(export(md, metadata={"translation_info": {"source_language": "de"}})).read_text()
        assert "Mustermann" not in content


def test_word_header_block_is_optional(tmp_path):
    meta = {"translation_info": {"source_language": "Deutsch", "target_language": "Englisch"}}
    plain = Document(EnhancedExportHandler(temp_dir=tmp_path).export_word("Text.", metadata=meta))
    assert [p.text for p in plain.paragraphs if p.text] == ["Text."]
    with_header = Document(EnhancedExportHandler(temp_dir=tmp_path, include_word_header=True)
                           .export_word("Text.", metadata=meta))
    assert len([p for p in with_header.paragraphs if p.text]) > 1


# --------------------------------------------------------------------------
# Markdown table helpers
# --------------------------------------------------------------------------
@pytest.mark.parametrize("line,expected", [
    ("|---|---|", True),
    ("| :-- | :-: | --: |", True),
    ("---|---", True),
    ("| | |", False),
    ("| a | b |", False),
    ("---", False),
])
def test_table_separator(line, expected):
    assert is_table_separator(line) is expected


def test_table_rows_and_validation():
    assert split_table_row("| a | b |") == ["a", "b"]
    assert count_table_columns("| a | b |") == 2
    assert is_valid_table_lines(["| a | b |", "|---|---|", "| 1 | 2 |"])
    assert not is_valid_table_lines(["| a | b |", "|---|---|", "| 1 |"])


# --------------------------------------------------------------------------
# LLM client: timeouts, retries, temperature
# --------------------------------------------------------------------------
def _run_with_server(handler, coro_factory):
    async def main():
        app = web.Application()
        app.router.add_post('/v1/chat/completions', handler)
        runner = web.AppRunner(app)
        await runner.setup()
        site = web.TCPSite(runner, '127.0.0.1', 0)
        await site.start()
        port = site._server.sockets[0].getsockname()[1]
        try:
            return await coro_factory(f"http://127.0.0.1:{port}/v1")
        finally:
            await runner.cleanup()
    return asyncio.run(main())


def _ok_response(content="ok"):
    return web.json_response({"choices": [{"message": {"content": content},
                                           "finish_reason": "stop"}], "model": "m"})


def test_timeout_is_the_configured_total(tmp_path):
    async def slow(request):
        await asyncio.sleep(2.0)
        return _ok_response()

    async def scenario(base):
        config = load_config()
        config.llm.api_base = base
        config.llm.timeout = 1
        config.llm.max_retries = 1
        client = LLMClient(config)
        await client.start()
        assert client._session.timeout.sock_read is None
        assert client._session.timeout.total == 1
        try:
            with pytest.raises(LLMTimeoutError) as err:
                await client.complete([Message(MessageRole.USER, "x")])
            assert "llm.timeout=1" in str(err.value)
        finally:
            await client.close()

    _run_with_server(slow, scenario)


def test_failing_requests_are_retried_once_per_layer():
    requests = []

    async def failing(request):
        requests.append(1)
        return web.Response(status=503)

    async def scenario(base):
        config = load_config()
        config.llm.api_base = base
        config.llm.max_retries = 3
        config.llm.retry_delay = 1
        config.llm.max_retry_delay = 1
        config.translation.auto_glossary = False
        client = LLMClient(config)
        await client.start()
        try:
            chunks = SimpleChunker(max_tokens=700, config=config).chunk_text("Ein Satz.", "de")
            engine = TableAwareTranslationEngine(client, config)
            result = await engine.translate_document(chunks, "de", "en", "preserve_original")
        finally:
            await client.close()
        return result

    result = _run_with_server(failing, scenario)
    # 3 attempts in the client, no additional engine-level retries.
    assert len(requests) == 3
    assert result.failed_chunks == [0]
    assert result.translated_chunks[0]['translated_content'] == "Ein Satz."


def test_temperature_only_sent_when_configured():
    bodies = []

    async def record(request):
        bodies.append(await request.json())
        return _ok_response()

    async def scenario(base):
        config = load_config()
        config.llm.api_base = base
        client = LLMClient(config)
        await client.start()
        try:
            config.llm.temperature = None
            await client.complete([Message(MessageRole.USER, "x")])
            config.llm.temperature = 0.4
            await client.complete([Message(MessageRole.USER, "x")])
        finally:
            await client.close()

    _run_with_server(record, scenario)
    assert "temperature" not in bodies[0]
    assert bodies[1]["temperature"] == 0.4


def test_no_atexit_registration_per_job():
    async def scenario():
        before = atexit._ncallbacks()
        for _ in range(5):
            client = LLMClient(load_config())
            await client.start()
            TableAwareTranslationEngine(client, client.config)
            await client.close()
        return atexit._ncallbacks() - before

    assert asyncio.run(scenario()) == 0


# --------------------------------------------------------------------------
# Translation engine
# --------------------------------------------------------------------------
def test_reasoning_blocks_are_removed():
    clean = TableAwareTranslationEngine._strip_reasoning
    assert clean("<think>plan</think>\n\nResult.") == "Result."
    assert clean("plan only</think>Result.") == "Result."
    assert clean("No reasoning.") == "No reasoning."


def test_identical_table_cells_are_translated_once():
    rows = ["| Measure | Owner | Status |", "|---|---|---|"]
    for i in range(10):
        rows.append(f"| Training part {i} | Personnel office | in progress |")
    text = "\n".join(rows)

    config = load_config()
    config.translation.auto_glossary = False
    client = EchoClient(config)
    chunks = SimpleChunker(max_tokens=2000, config=config).chunk_text(text, "en")
    engine = TableAwareTranslationEngine(client, config)
    result = asyncio.run(engine.translate_document(chunks, "en", "de", "preserve_original"))

    # 3 headers + 10 distinct first cells + 1 + 1 repeated values = 15 calls
    assert len(client.calls) == 15
    assert result.translated_chunks[0]['translated_content'].count("Personnel office") == 10


def test_glossary_candidates_respect_cells_and_sentence_starts():
    text = ("# Projektplan\n\n"
            "Weitere Schritte folgen. Die Projektgruppe prüft den Projektplan.\n"
            "Danach informiert die Projektgruppe die Leitung.\n\n"
            "| Aufgabe | Status |\n|---|---|\n"
            "| Schulung | in Bearbeitung |\n| Schulung | Termin offen |\n")
    config = load_config()
    chunks = SimpleChunker(max_tokens=2000, config=config).chunk_text(text, "de")
    engine = TableAwareTranslationEngine(EchoClient(config), config)
    terms = engine._extract_term_candidates(chunks.chunks, {}, 100)

    assert "Projektgruppe" in terms and "Projektplan" in terms and "Schulung" in terms
    assert "Weitere Schritte" not in terms          # sentence start
    assert not any("Bearbeitung" in t and " " in t for t in terms)  # no cross-cell join
    assert "Leitung" not in terms                   # occurs once only


def test_local_structural_check_needs_no_model():
    config = load_config()
    client = EchoClient(config)
    engine = TableAwareTranslationEngine(client, config)
    chunks = [{'original_content': "Es kostet 1500 Euro.",
               'translated_content': "It costs 1600 euros.", 'position': 0}]
    findings = asyncio.run(engine.verify_translation(chunks, "de", "en", semantic=False))
    assert findings and client.calls == []


# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------
def test_config_is_independent_of_working_directory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    config = load_config()
    assert config.simple_chunking.max_tokens == 700  # value from config.yaml
    assert not (tmp_path / "temp").exists()


def test_local_overrides_and_environment(tmp_path, monkeypatch):
    local = tmp_path / "config.local.yaml"
    local.write_text("llm:\n  model: local-model\ninterface:\n  host: 0.0.0.0\n")
    monkeypatch.setenv("TRANSLATION_LOGGING__LEVEL", "DEBUG")
    config = load_config(local_config_path=local)
    assert config.llm.model == "local-model"
    assert config.llm.api_base == "http://localhost:11434/v1"  # untouched
    assert config.interface.host == "0.0.0.0"
    assert config.logging.level == "DEBUG"


def test_deep_merge_and_resolve_path():
    merged = deep_merge({"a": {"b": 1, "c": 2}}, {"a": {"c": 3}})
    assert merged == {"a": {"b": 1, "c": 3}}
    assert resolve_path("temp") == PROJECT_ROOT / "temp"
    assert resolve_path("/abs/path") == Path("/abs/path")


def test_safe_defaults():
    config = load_config()
    assert config.interface.host == "127.0.0.1"
    assert config.llm.temperature is None
    assert config.export.word.include_header is False
    assert config.interface.footer_links == []


def test_log_file_is_written(tmp_path, monkeypatch):
    log_file = tmp_path / "app.log"
    monkeypatch.setenv("TRANSLATION_LOGGING__FILE", str(log_file))
    logging.disable(logging.NOTSET)
    try:
        load_config()
        logging.getLogger("test").info("hello")
        for handler in logging.getLogger().handlers:
            handler.flush()
    finally:
        logging.disable(logging.CRITICAL)
    assert log_file.exists() and "hello" in log_file.read_text()


# --------------------------------------------------------------------------
# Application
# --------------------------------------------------------------------------
def test_style_keys_drive_formality():
    from src.app import style_label, style_key_for_label, formality_for_style
    assert style_key_for_label(style_label("professionell_du")) == "professionell_du"
    assert formality_for_style("professionell_du") == "informal"
    assert formality_for_style("professionell_sie") == "formal"
    assert style_key_for_label(style_label("custom_style")) == "custom_style"


def test_stale_temp_dirs_are_removed(tmp_path):
    from src.app import cleanup_stale_temp_dirs, OWNER_FILE
    live = tmp_path / "tmplive"
    live.mkdir()
    (live / OWNER_FILE).write_text(str(os.getpid()))
    dead = tmp_path / "tmpdead"
    dead.mkdir()
    (dead / OWNER_FILE).write_text("999999999")
    old = tmp_path / "tmpold"
    old.mkdir()
    os.utime(old, (time.time() - 7200, time.time() - 7200))
    fresh = tmp_path / "tmpfresh"
    fresh.mkdir()

    assert cleanup_stale_temp_dirs(tmp_path) == 2
    assert live.exists() and fresh.exists()
    assert not dead.exists() and not old.exists()


def test_interface_exposes_no_api_endpoints():
    from src.app import UnifiedDocumentTranslationApp
    app = UnifiedDocumentTranslationApp()
    try:
        interface = app.create_unified_interface()
        visibilities = {fn.api_visibility for fn in interface.fns.values()}
        assert visibilities == {"private"}
        assert app.markdown_converter.include_frontmatter is False
    finally:
        app.cleanup()


def test_pdf_ocr_settings_are_passed(monkeypatch):
    import pymupdf4llm
    from src.document_processing.pdf_processor import PDFProcessor
    if not getattr(pymupdf4llm, '_use_layout', False):
        pytest.skip("pymupdf4llm layout path (with OCR options) not installed")
    seen = {}
    monkeypatch.setattr(pymupdf4llm, "to_markdown", lambda path, **kw: seen.update(kw) or "x")
    config = load_config()
    config.document_processing.pdf_ocr = False
    config.document_processing.pdf_ocr_language = "deu"
    PDFProcessor(config=config)._to_markdown(Path("dummy.pdf"))
    assert seen["use_ocr"] is False and seen["ocr_language"] == "deu"


def test_app_shares_its_configuration():
    from src.app import UnifiedDocumentTranslationApp
    from src.config import get_config
    app = UnifiedDocumentTranslationApp()
    try:
        assert get_config() is app.config
    finally:
        app.cleanup()


def test_upload_cache_is_limited():
    from src.app import UnifiedDocumentTranslationApp
    app = UnifiedDocumentTranslationApp()
    try:
        app.config.document_processing.upload_retention_minutes = 60
        assert app._upload_cache_policy() == (600, 3600)
        interface = app.create_unified_interface()
        assert interface.delete_cache == (600, 3600)
        app.config.document_processing.upload_retention_minutes = None
        assert app._upload_cache_policy() is None
    finally:
        app.cleanup()


# --------------------------------------------------------------------------
# input files and error statuses
# --------------------------------------------------------------------------
class _Upload:
    def __init__(self, name):
        self.name = name


def _run_app(app, kind, text="", path=None):
    return asyncio.run(app.process_and_translate_unified(
        kind, text, _Upload(str(path)) if path else None, "de", "en", "direct", "", "",
        progress=lambda *a, **k: None, lang="en"))


@pytest.fixture
def app_with_echo(monkeypatch):
    from src.app import UnifiedDocumentTranslationApp
    import src.app as app_module

    class Echo(EchoClient):
        def __init__(self, config, shared_semaphore=None):
            super().__init__(config)

    monkeypatch.setattr(app_module, "LLMClient", Echo)
    app = UnifiedDocumentTranslationApp()
    app.config.translation.auto_glossary = False
    yield app
    app.cleanup()


def test_markdown_file_is_used_verbatim(tmp_path, app_with_echo):
    md = "# Title\n\n| A | B |\n|---|---|\n| 1 | x |\n| 2 | y |\n\n1. one\n2. two\n\n```\n# not a heading\ncode()\n```\n"
    path = tmp_path / "doc.md"
    path.write_text(md)
    content, translation, status, _ = _run_app(app_with_echo, "File", path=path)
    assert content == md
    assert "| 1 | x |\n| 2 | y |" in translation and "1. one\n2. two" in translation
    assert "tables" in status


@pytest.mark.parametrize("data", [
    "Größe und Ärger".encode("latin-1"),
    "„Zitat“ und €".encode("cp1252"),
    "\ufeffMit BOM: Größe".encode("utf-8"),
    "UTF-16: Größe".encode("utf-16"),
])
def test_text_files_are_decoded_independently_of_locale(tmp_path, data):
    from src.document_processing import decode_text, process_document
    path = tmp_path / "t.txt"
    path.write_bytes(data)
    result = asyncio.run(process_document(path))
    assert result.success
    assert result.document.raw_text == decode_text(data)
    assert "\\x" not in result.document.raw_text and not result.document.raw_text.startswith("\ufeff")


def test_blank_text_asks_for_input(app_with_echo):
    _, translation, status, _ = _run_app(app_with_echo, "Text", text="   \n ")
    assert translation == "" and status == "Please enter a text."


def test_unreadable_file_message_without_server_paths(tmp_path, app_with_echo):
    path = tmp_path / "broken.docx"
    path.write_text("not a word file")
    _, translation, status, _ = _run_app(app_with_echo, "File", path=path)
    assert translation == "" and "broken.docx could not be read" in status
    assert str(app_with_echo.temp_dir) not in status


def test_total_failure_is_reported_as_failure(monkeypatch):
    from src.app import UnifiedDocumentTranslationApp
    app = UnifiedDocumentTranslationApp()
    try:
        app.config.llm.api_base = "http://127.0.0.1:1/v1"
        app.config.llm.max_retries = 1
        app.config.translation.auto_glossary = False
        _, translation, status, _ = _run_app(app, "Text", text="Ein Satz.")
        assert translation == "" and status.startswith("❌ No section could be translated")
    finally:
        app.cleanup()


def test_document_conversion_runs_outside_the_event_loop(tmp_path, monkeypatch):
    import threading
    import src.app as app_module
    seen = {}

    async def fake_process(path, options=None):
        seen["thread"] = threading.current_thread().name
        return "ok"

    monkeypatch.setattr(app_module, "process_document", fake_process)

    async def scenario():
        main = threading.current_thread().name
        for name in ("x.pdf", "x.docx"):
            assert await app_module.convert_document_off_loop(tmp_path / name) == "ok"
            assert seen["thread"] != main
        return seen["thread"]

    asyncio.run(scenario())


def test_html_export_escapes_raw_html(tmp_path):
    html = Path(EnhancedExportHandler(temp_dir=tmp_path).export_html(
        "Text <script>alert(1)</script> und a < b.\n\n| A | B |\n|---|---|\n| 1 | 2 |")).read_text()
    assert "<script>alert" not in html and "&lt;script&gt;" in html
    assert "<table" in html


def test_invalid_target_language_is_rejected(app_with_echo):
    _, translation, status, _ = asyncio.run(app_with_echo.process_and_translate_unified(
        "Text", "Ein Satz.", None, "de", "auto", "direct", "", "",
        progress=lambda *a, **k: None, lang="en"))
    assert translation == "" and status == "❌ Please choose a target language."
