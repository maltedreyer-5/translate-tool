# SPDX-FileCopyrightText: 2026 Malte Dreyer
# SPDX-License-Identifier: MIT
"""
Answer length limits: per-request budget, retry with the full window,
splitting of cut-off sections and of over-long Word paragraphs.
"""
import asyncio
import logging
import re

import pytest
from docx import Document

from src.config import load_config
from src.llm_client import (LLMClient, CompletionResponse, Message, MessageRole,
                            LLMTruncatedError)
from src.chunking.simple_chunker import SimpleChunker
from src.translation.translation_engine import TableAwareTranslationEngine, split_text_in_two
from src.translation.docx_inplace import DocxInPlaceTranslator

logging.disable(logging.CRITICAL)


def upper_keep_markers(text):
    return re.sub(r'(<[^>]+>)|([^<]+)', lambda m: m.group(1) or m.group(2).upper(), text)


class LimitedClient(LLMClient):
    """
    Answers like a model that can produce at most `capacity` characters per
    answer, and that is cut off when max_tokens is smaller than the answer
    (2 characters per token).
    """

    def __init__(self, config, capacity=10**9):
        super().__init__(config)
        self.capacity = capacity
        self.calls = []

    async def complete(self, messages, temperature=None, max_tokens=None, stream=False, **kw):
        # Exercise the real budget logic of LLMClient.complete() ...
        return await super().complete(messages, temperature, max_tokens, stream, **kw)

    async def _complete_once(self, messages, temperature, max_tokens, stream, **kw):
        # ... but answer locally instead of over HTTP.
        user = messages[-1].content
        text = user.split("\n\n", 1)[1] if "\n\n" in user else user
        answer = upper_keep_markers(text)
        if max_tokens is None:  # as in the real _complete_once: fixed config value
            max_tokens = self.llm_config.max_tokens
        self.calls.append({"chars": len(text), "max_tokens": max_tokens})
        limit = min(self.capacity, max_tokens * 2)
        if len(answer) > limit:
            return CompletionResponse(content=answer[:limit], model="m", finish_reason="length")
        return CompletionResponse(content=answer, model="m", finish_reason="stop")


def cfg(**llm):
    c = load_config()
    c.translation.auto_glossary = False
    for k, v in llm.items():
        setattr(c.llm, k, v)
    return c


# --------------------------------------------------------------------------
# budget
# --------------------------------------------------------------------------
def test_budget_grows_with_the_text_and_respects_the_window():
    client = LLMClient(cfg(context_window=32000))
    short, cap = client._output_budget([Message(MessageRole.USER, "x" * 100)])
    long, cap2 = client._output_budget([Message(MessageRole.USER, "x" * 6000)])
    assert 256 <= short < long <= cap2
    huge, cap3 = client._output_budget([Message(MessageRole.SYSTEM, "s" * 40000),
                                        Message(MessageRole.USER, "x" * 30000)])
    assert huge == cap3 and cap3 <= 32000 - 40000 // 3


def test_reasoning_reserve_is_added():
    plain = LLMClient(cfg())._output_budget([Message(MessageRole.USER, "x" * 100)])[0]
    reserve = LLMClient(cfg(reasoning_reserve_tokens=2000))._output_budget(
        [Message(MessageRole.USER, "x" * 100)])[0]
    assert reserve == plain + 2000


def test_fixed_limit_is_sent_unchanged():
    client = LimitedClient(cfg(max_tokens=1234))
    asyncio.run(client.complete([Message(MessageRole.USER, "Hallo")]))
    assert client.calls[0]["max_tokens"] == 1234


def test_cut_off_answer_is_requested_again_with_the_full_window():
    client = LimitedClient(cfg(reasoning_reserve_tokens=0))
    client._output_budget = lambda messages: (50, 5000)   # estimate too small
    response = asyncio.run(client.complete([Message(MessageRole.USER, "x" * 400)]))
    assert [c["max_tokens"] for c in client.calls] == [50, 5000]
    assert response.finish_reason == "stop"


def test_translate_text_reports_truncation_instead_of_returning_partial_text():
    client = LimitedClient(cfg(), capacity=100)
    with pytest.raises(LLMTruncatedError):
        asyncio.run(client.translate_text("Wort " * 100, "de", "en"))


# --------------------------------------------------------------------------
# splitting in the text pipeline
# --------------------------------------------------------------------------
def test_split_prefers_paragraphs_then_sentences():
    first, sep, second = split_text_in_two("A a a.\n\nB b b.")
    assert (first, sep, second) == ("A a a.", "\n\n", "B b b.")
    first, sep, second = split_text_in_two("Eins zwei. Drei vier. Fünf sechs.")
    assert sep == " " and first.endswith(".")
    assert split_text_in_two("unteilbar") is None


def test_cut_off_section_is_split_and_nothing_is_lost():
    c = cfg()
    text = " ".join(f"Satz {i} hat einige Wörter." for i in range(120))   # ~3000 chars
    client = LimitedClient(c, capacity=900)       # the "model" answers at most 900 chars
    engine = TableAwareTranslationEngine(client, c)
    chunks = SimpleChunker(max_tokens=2000, config=c).chunk_text(text, "de")
    result = asyncio.run(engine.translate_document(chunks, "de", "en", "preserve_original"))
    assert result.failed_chunks == []
    assert result.translated_chunks[0]['translated_content'] == text.upper()
    assert engine.stats.get('split_sections', 0) >= 2


# --------------------------------------------------------------------------
# long Word paragraphs
# --------------------------------------------------------------------------
def _translate_docx(tmp_path, build, source_lang="de", max_tokens=200):
    c = cfg()
    c.simple_chunking.max_tokens = max_tokens
    doc = Document()
    build(doc)
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    doc.save(src)
    client = LimitedClient(c)
    engine = TableAwareTranslationEngine(client, c)
    translator = DocxInPlaceTranslator(engine, c)
    result = asyncio.run(translator.translate_file(src, out, source_lang, "en",
                                                   "preserve_original", "formal", {}))
    return Document(out), result, client, translator


def test_long_plain_paragraph_is_sent_in_pieces(tmp_path):
    text = " ".join(f"Satz {i} der langen Darstellung." for i in range(150))  # ~4800 chars
    doc, result, client, translator = _translate_docx(tmp_path, lambda d: d.add_paragraph(text))
    assert doc.paragraphs[0].text == text.upper()
    assert result.paragraphs == 1 and result.failed_paragraphs == 0
    assert max(c["chars"] for c in client.calls) <= translator.max_chars + 20
    assert len(client.calls) >= 6


def test_long_formatted_paragraph_keeps_its_formatting(tmp_path):
    before = " ".join(f"Vorne {i} steht Text." for i in range(70))
    after = " ".join(f"Hinten {i} folgt mehr." for i in range(70))

    def build(d):
        p = d.add_paragraph(before + " ")
        p.add_run("FETTES WORT").bold = True
        p.add_run(" " + after)
    doc, result, client, translator = _translate_docx(tmp_path, build)
    paragraph = doc.paragraphs[0]
    assert paragraph.text == (before + " FETTES WORT " + after).upper()
    assert [r.text for r in paragraph.runs if r.bold] == ["FETTES WORT"]
    assert result.formatting_fallbacks == 0
    assert max(c["chars"] for c in client.calls) <= translator.max_chars + 40


def test_section_size_follows_the_source_language(tmp_path):
    _, _, _, german = _translate_docx(tmp_path, lambda d: d.add_paragraph("Text."), "de", 700)
    _, _, _, chinese = _translate_docx(tmp_path, lambda d: d.add_paragraph("文本。"), "zh", 700)
    assert german.max_chars == int(700 * 3.5)
    assert chinese.max_chars == int(700 * 1.5)
